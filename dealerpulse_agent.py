"""
DealerPulse AI — Cars24 Dealer Activity & Performance Intelligence Agent
TM: rishabh.malhotra1@cars24.com

Usage:
    python3 dealerpulse_agent.py <path_to_dealer_dump.csv>
    python3 dealerpulse_agent.py <path_to_dealer_dump.csv> --excel
    python3 dealerpulse_agent.py <path_to_dealer_dump.csv> --slack-preview
"""

import sys, csv, json, math, os
from datetime import datetime
from collections import defaultdict

TARGET_TM   = "rishabh.malhotra1@cars24.com"
LAST3       = [12, 13, 14]          # Feb-26, Mar-26, Apr-26
LAST6       = list(range(9, 15))    # Nov-25 to Apr-26
HIST        = list(range(2, 12))    # Apr-25 to Jan-26
MONTH_NAMES = {
    2:"Apr-25",3:"May-25",4:"Jun-25",5:"Jul-25",6:"Aug-25",7:"Sep-25",
    8:"Oct-25",9:"Nov-25",10:"Dec-25",11:"Jan-26",12:"Feb-26",13:"Mar-26",14:"Apr-26"
}

# ── CSV Parser ─────────────────────────────────────────────────────────────
def safe_int(val):
    try: return int(float(str(val).strip())) if str(val).strip() not in ("", "None", "nan") else 0
    except: return 0

def load_dealers(csv_path):
    dealers = []
    with open(csv_path, "r", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        headers = next(reader)
        for row in reader:
            if len(row) < 84: continue
            tl = row[81].strip().lower()
            tm = row[82].strip().lower()
            if TARGET_TM not in (tl, tm): continue

            leads = [safe_int(row[6 + i]) for i in range(15)]
            insps = [safe_int(row[26 + i]) for i in range(15)]
            sis   = [safe_int(row[44 + i]) for i in range(15)]

            l3_leads = sum(leads[i] for i in LAST3)
            l3_insp  = sum(insps[i] for i in LAST3)
            l3_si    = sum(sis[i]   for i in LAST3)
            l6_leads = sum(leads[i] for i in LAST6)
            hist_leads = sum(leads[i] for i in HIST)

            # Categorise
            if l3_leads > 0 and l3_insp > 0 and l3_si > 0:
                cat = "TOP"
                if l3_si >= 5:   tier = "GOLD"
                elif l3_si >= 2: tier = "DIAMOND"
                else:            tier = "SILVER"
            elif l6_leads == 0:
                cat, tier = "DORMANT", "DORMANT"
            elif l3_leads == 0 and (hist_leads > 0 or l6_leads > 0):
                cat, tier = "CHURNED", "CHURNED"
            elif l3_leads > 0:
                cat, tier = "ACTIVE-NO-SI", "ACTIVE-NO-SI"
            else:
                cat, tier = "DORMANT", "DORMANT"

            total_leads = safe_int(row[22])
            total_insp  = safe_int(row[42])
            total_si    = safe_int(row[60])
            a2i  = round(l3_insp / l3_leads * 100, 1)  if l3_leads  > 0 else 0
            i2si = round(l3_si   / l3_insp  * 100, 1)  if l3_insp   > 0 else 0

            # Month-wise leads trend
            monthly = {MONTH_NAMES[i]: leads[i] for i in sorted(MONTH_NAMES.keys())}

            dealers.append({
                "code":        row[0].strip(),
                "name":        row[1].strip(),
                "region":      row[2].strip(),
                "zone":        row[3].strip(),
                "kam":         row[5].strip(),
                "tl":          tl,
                "tm":          tm,
                "onboard":     row[83].strip() if len(row) > 83 else "",
                "category":    cat,
                "tier":        tier,
                "l3_leads":    l3_leads,
                "l3_insp":     l3_insp,
                "l3_si":       l3_si,
                "l6_leads":    l6_leads,
                "hist_leads":  hist_leads,
                "total_leads": total_leads,
                "total_insp":  total_insp,
                "total_si":    total_si,
                "a2i":         a2i,
                "i2si":        i2si,
                "monthly":     monthly,
            })
    return dealers

# ── Analytics ──────────────────────────────────────────────────────────────
def analyse(dealers):
    counts = defaultdict(int)
    for d in dealers:
        counts[d["tier"]] += 1

    kam_map = defaultdict(list)
    for d in dealers:
        kam_map[d["kam"]].append(d)

    kam_scores = {}
    for kam, ds in kam_map.items():
        total_si = sum(d["l3_si"] for d in ds)
        active   = sum(1 for d in ds if d["l3_leads"] > 0)
        churned  = sum(1 for d in ds if d["tier"] == "CHURNED")
        gold     = sum(1 for d in ds if d["tier"] == "GOLD")
        diamond  = sum(1 for d in ds if d["tier"] == "DIAMOND")
        silver   = sum(1 for d in ds if d["tier"] == "SILVER")
        if total_si >= 90:   grade = "🟢 STAR"
        elif total_si >= 50: grade = "🟡 GOOD"
        elif total_si >= 20: grade = "🟠 AVERAGE"
        else:                grade = "🔴 NEEDS ACTION"
        kam_scores[kam] = {
            "total": len(ds), "active": active, "churned": churned,
            "gold": gold, "diamond": diamond, "silver": silver,
            "l3_si": total_si, "grade": grade
        }

    churn_risk = sorted(
        [d for d in dealers if d["tier"] == "CHURNED"],
        key=lambda x: x["hist_leads"], reverse=True
    )[:10]

    top_dealers = sorted(
        [d for d in dealers if d["tier"] in ("GOLD","DIAMOND","SILVER")],
        key=lambda x: x["l3_si"], reverse=True
    )[:10]

    return {"counts": dict(counts), "kam": kam_scores, "churn_risk": churn_risk, "top": top_dealers}

# ── Slack Flash Card ────────────────────────────────────────────────────────
def build_slack_blocks(dealers, stats, date_str):
    counts  = stats["counts"]
    total   = len(dealers)
    gold    = counts.get("GOLD", 0)
    diamond = counts.get("DIAMOND", 0)
    silver  = counts.get("SILVER", 0)
    active  = counts.get("ACTIVE-NO-SI", 0)
    churned = counts.get("CHURNED", 0)
    dormant = counts.get("DORMANT", 0)

    top_text = "\n".join(
        f"{'🥇' if d['tier']=='GOLD' else '💎' if d['tier']=='DIAMOND' else '🥈'} "
        f"*{d['name'][:22]}* — {d['l3_si']} SI | {d['l3_leads']} Leads | A2I {d['a2i']}%"
        for d in stats["top"][:5]
    ) or "_No top dealers_"

    churn_text = "\n".join(
        f"🔴 *{d['name'][:22]}* ({d['region']}) — Last active: {max((k for k,v in d['monthly'].items() if v>0), default='N/A')} | KAM: {d['kam'][:15]}"
        for d in stats["churn_risk"][:5]
    ) or "_No churn alerts_"

    kam_lines = sorted(stats["kam"].items(), key=lambda x: x[1]["l3_si"], reverse=True)
    kam_text = "\n".join(
        f"{v['grade']} *{k[:18]}* — {v['total']} dealers | 🥇{v['gold']} 💎{v['diamond']} 🥈{v['silver']} | {v['l3_si']} SI | Churned: {v['churned']}"
        for k, v in kam_lines[:8]
    ) or "_No KAM data_"

    blocks = [
        {"type": "header", "text": {"type": "plain_text", "text": f"📊 DealerPulse AI — Daily Flash Card | {date_str}"}},
        {"type": "section", "text": {"type": "mrkdwn",
            "text": f"*Network Snapshot — {total} Dealers*\n"
                    f"🥇 Gold: *{gold}*  💎 Diamond: *{diamond}*  🥈 Silver: *{silver}*  "
                    f"⚠️ Active-No-SI: *{active}*  🔴 Churned: *{churned}*  💤 Dormant: *{dormant}*"}},
        {"type": "divider"},
        {"type": "section", "text": {"type": "mrkdwn", "text": f"*🏆 Top Performers (L3M)*\n{top_text}"}},
        {"type": "divider"},
        {"type": "section", "text": {"type": "mrkdwn", "text": f"*🚨 Churn Risk — Act Today*\n{churn_text}"}},
        {"type": "divider"},
        {"type": "section", "text": {"type": "mrkdwn", "text": f"*👤 KAM Scorecards*\n{kam_text}"}},
        {"type": "context", "elements": [{"type": "mrkdwn", "text": f"_DealerPulse AI | TM: {TARGET_TM} | Generated {datetime.now().strftime('%d %b %Y %H:%M IST')}_"}]}
    ]
    return blocks

# ── Excel Export ───────────────────────────────────────────────────────────
def export_excel(dealers, out_path):
    try:
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    except ImportError:
        print("Installing openpyxl..."); os.system("pip install openpyxl -q")
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

    wb = openpyxl.Workbook()

    TIER_ORDER  = ["GOLD","DIAMOND","SILVER","ACTIVE-NO-SI","CHURNED","DORMANT"]
    TIER_COLORS = {"GOLD":"FFD700","DIAMOND":"B9F2FF","SILVER":"C0C0C0",
                   "ACTIVE-NO-SI":"FFF3CD","CHURNED":"FFCCCC","DORMANT":"E8E8E8"}
    TIER_ICONS  = {"GOLD":"🥇 GOLD","DIAMOND":"💎 DIAMOND","SILVER":"🥈 SILVER",
                   "ACTIVE-NO-SI":"⚠ Active-No-SI","CHURNED":"🔴 Churned","DORMANT":"💤 Dormant"}
    ACTION_MAP  = {
        "GOLD":        "Retain & upsell — protect relationship",
        "DIAMOND":     "Grow to Gold — push for 5+ SI",
        "SILVER":      "Nurture — help reach 2 SI",
        "ACTIVE-NO-SI":"Convert inspections to stock-in",
        "CHURNED":     "Re-engage — call immediately",
        "DORMANT":     "Win-back campaign",
    }
    PRIORITY_MAP = {"GOLD":"P1","DIAMOND":"P1","SILVER":"P2","ACTIVE-NO-SI":"P2","CHURNED":"P1","DORMANT":"P3"}

    COLS = ["Dealer Code","Dealer Name","Region","Zone","KAM","Tier",
            "Onboard Date","L3M Leads","L3M Insp","L3M SI","A2I%","I2SI%",
            "Total Leads","Total Insp","Total SI","Action Required","Priority"]

    def hdr_fill(hex_col):
        return PatternFill("solid", fgColor=hex_col)

    def write_sheet(ws, rows, title_color):
        ws.append(COLS)
        for cell in ws[1]:
            cell.font      = Font(bold=True, color="FFFFFF")
            cell.fill      = hdr_fill(title_color)
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            action   = ACTION_MAP.get(row["tier"], "")
            priority = PRIORITY_MAP.get(row["tier"], "P4")
            ws.append([
                row["code"], row["name"], row["region"], row["zone"], row["kam"],
                TIER_ICONS.get(row["tier"], row["tier"]),
                row["onboard"], row["l3_leads"], row["l3_insp"], row["l3_si"],
                row["a2i"], row["i2si"], row["total_leads"], row["total_insp"],
                row["total_si"], action, priority
            ])
            fill_hex = TIER_COLORS.get(row["tier"], "FFFFFF")
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=fill_hex)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(
                len(str(c.value or "")) for c in col) + 3
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    # Summary sheet
    ws0 = wb.active
    ws0.title = "📊 Summary"
    counts = defaultdict(int)
    for d in dealers: counts[d["tier"]] += 1
    ws0.append(["DealerPulse AI — Portfolio Summary"])
    ws0.append([f"Generated: {datetime.now().strftime('%d %b %Y %H:%M')}"])
    ws0.append([f"TM: {TARGET_TM}"])
    ws0.append([])
    ws0.append(["Tier","Count","% of Portfolio"])
    total = len(dealers)
    for t in TIER_ORDER:
        c = counts.get(t, 0)
        ws0.append([TIER_ICONS[t], c, f"{c/total*100:.1f}%"])
    ws0.append(["TOTAL", total, "100%"])

    # Tier sheets
    COLOR_MAP = {"GOLD":"B8860B","DIAMOND":"0070C0","SILVER":"696969",
                 "ACTIVE-NO-SI":"C55A11","CHURNED":"C00000","DORMANT":"595959"}
    for tier in TIER_ORDER:
        rows = [d for d in dealers if d["tier"] == tier]
        ws   = wb.create_sheet(TIER_ICONS[tier])
        write_sheet(ws, rows, COLOR_MAP[tier])

    # KAM sheet
    ws_kam = wb.create_sheet("📋 KAM-wise")
    write_sheet(ws_kam, sorted(dealers, key=lambda x: (x["kam"], -x["l3_si"])), "3A3A3A")

    wb.save(out_path)
    print(f"✅ Excel saved → {out_path}  ({os.path.getsize(out_path)//1024} KB)")

# ── n8n Payload Generator ───────────────────────────────────────────────────
def save_n8n_payload(dealers, stats, out_dir):
    date_str = datetime.now().strftime("%d %b %Y")
    blocks   = build_slack_blocks(dealers, stats, date_str)
    payload  = {
        "generated_at": datetime.now().isoformat(),
        "date_str":     date_str,
        "total_dealers": len(dealers),
        "counts":       stats["counts"],
        "slack_blocks": blocks,
    }
    path = os.path.join(out_dir, "latest_flash_card.json")
    with open(path, "w") as f:
        json.dump(payload, f, indent=2)
    print(f"✅ n8n payload → {path}")
    return path

# ── MAIN ───────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python3 dealerpulse_agent.py <dump.csv> [--excel] [--slack-preview]")
        sys.exit(1)

    csv_path = sys.argv[1]
    do_excel = "--excel" in sys.argv or "--all" in sys.argv
    do_slack = "--slack-preview" in sys.argv or "--all" in sys.argv

    print(f"\n{'='*60}")
    print(f"  DealerPulse AI  |  {datetime.now().strftime('%d %b %Y %H:%M')}")
    print(f"  TM: {TARGET_TM}")
    print(f"{'='*60}\n")

    print("📂 Loading dealer data...")
    dealers = load_dealers(csv_path)
    print(f"✅ {len(dealers)} dealers loaded (filtered to TM/TL = {TARGET_TM})\n")

    print("🔍 Analysing...")
    stats  = analyse(dealers)
    counts = stats["counts"]

    print(f"  🥇 GOLD        : {counts.get('GOLD',0)}")
    print(f"  💎 DIAMOND     : {counts.get('DIAMOND',0)}")
    print(f"  🥈 SILVER      : {counts.get('SILVER',0)}")
    print(f"  ⚠  Active-No-SI: {counts.get('ACTIVE-NO-SI',0)}")
    print(f"  🔴 Churned     : {counts.get('CHURNED',0)}")
    print(f"  💤 Dormant     : {counts.get('DORMANT',0)}")
    print(f"  📊 Total       : {len(dealers)}\n")

    out_dir = os.path.dirname(os.path.abspath(csv_path))

    # Always save n8n payload
    save_n8n_payload(dealers, stats, "/home/user/Dealer-Pulse-AI")

    if do_excel:
        xl_path = os.path.join("/home/user/Dealer-Pulse-AI", "DealerPulse_Portfolio_Report.xlsx")
        print("\n📊 Generating Excel report...")
        export_excel(dealers, xl_path)

    if do_slack:
        print("\n📨 Slack preview (blocks JSON):")
        blocks = build_slack_blocks(dealers, stats, datetime.now().strftime("%d %b %Y"))
        print(json.dumps(blocks, indent=2)[:2000], "...\n")

    print("\n✅ Done. Drop latest_flash_card.json into n8n to publish to Slack.")
    print("   Or run with --excel to regenerate the full portfolio Excel.\n")
